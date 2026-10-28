"""Campaign sending from the app's Mailing panel: plan, run in the background, stop, log.

Same rules as `python -m mailing.send --send`: today's limit per mailbox, mailboxes take
turns, nobody gets it twice (envios.csv), unsubscribes and bounces are checked first.
Nothing goes out unless the request carries confirm="ENVIAR", typed by the person.
"""
import csv
import io
import re
import subprocess
import sys
import threading
from datetime import datetime
from pathlib import Path

from . import deliverability, library, replies, send
from .build import load_campaign

CONFIRM_WORD = "ENVIAR"
RECENT_ROWS = 40


def campaign_path(out_dir, slug=""):
    """Folder of a campaign of this person's library (<out_dir>/campanas); with no valid
    name, the one chosen in the Campaña tab."""
    return library.pick(out_dir, slug)


def paths(out_dir, campaign):
    out_dir = Path(out_dir)
    return {"contacts": out_dir / "contacts.csv", "out": out_dir / "mailing" / campaign["slug"],
            "suppression": out_dir / "suppression.txt"}


def available_senders(env_path, campaign):
    from mail_contact_scraper import accounts, envfile

    found = [m["email"] for m in accounts.list_mailboxes(env_path) if m["has_password"]]
    default = campaign["sender"]["email"]
    if envfile.environ(env_path).get("SMTP_USER") and default.lower() not in (e.lower() for e in found):
        found.insert(0, default)  # a relay can send as any address of the domain
    return found


def _senders(env_path, campaign, chosen, daily_limit):
    available = available_senders(env_path, campaign)
    lowered = {e.lower(): e for e in available}
    chosen = [lowered[c.lower()] for c in (chosen or []) if c and c.lower() in lowered]
    if not chosen:
        default = campaign["sender"]["email"].lower()
        chosen = [lowered[default]] if default in lowered else available[:1]
    if not chosen:
        raise ValueError("Primero agrega un buzón con su contraseña de aplicación en Configuración ▸ Buzones.")
    try:
        return send.load_senders(chosen, campaign, env_path, daily_limit or None)
    except SystemExit as exc:
        raise ValueError(str(exc)) from None


def _limit(value, name):
    try:
        number = int(value or 0)
    except (TypeError, ValueError):
        raise ValueError(f"{name} tiene que ser un número.") from None
    if number < 0:
        raise ValueError(f"{name} no puede ser negativo.")
    return number


def _per_company(value):
    return send.PER_COMPANY_DAILY if value in (None, "") else _limit(value, "El máximo por empresa")


def overview(out_dir, env_path, run, slug="", chosen=(), daily_limit=0, per_company=None):
    """Everything the Campaña tab shows: list, today's room per mailbox, log, run, and the
    follow-up (who is due, who answered)."""
    path = campaign_path(out_dir, slug)
    campaign = load_campaign(path)
    where = paths(out_dir, campaign)
    log = send.read_log(where["out"] / send.LOG_NAME)
    counts = {}
    for row in log:
        counts[row["estado"]] = counts.get(row["estado"], 0) + 1
    info = {
        "campaign": path.stem, "name": campaign["name"], "subject": campaign["subject"],
        "available": available_senders(env_path, campaign),
        "default_sender": campaign["sender"]["email"], "plan": None, "error": "",
        "log": {"counts": counts, "recent": [
            {k: row.get(k, "") for k in ("fecha", "email", "estado", "detalle", "remitente")}
            for row in reversed(log[-RECENT_ROWS:])]},
        "suppression": send.suppression_counts(where["suppression"]),
        "run": run.status(),
        "replies": _reply_counts(where["out"]),
        "followup": None,
        "followup_text": _followup_text(campaign),
        "per_company": send.PER_COMPANY_DAILY,
    }
    try:
        senders, _ = _senders(env_path, campaign, chosen, _limit(daily_limit, "El límite diario"))
    except ValueError as exc:
        info["error"] = str(exc)
        return info
    info["senders"] = [{"email": s["email"], "daily_limit": s["daily_limit"]} for s in senders]
    if not where["contacts"].exists():
        info["error"] = "Todavía no hay contactos. Corre un escaneo primero."
        return info
    per_company = _per_company(per_company)
    plan = send.send_campaign(path, where["contacts"], where["out"], senders=senders, per_company=per_company,
                              suppression_path=where["suppression"], say=lambda *a: None)
    keys = ("audience", "already_sent", "pending", "sent_today", "room", "this_run", "skipped", "held_by_company")
    info["plan"] = {k: plan[k] for k in keys}
    if campaign.get("followup", {}).get("body"):
        due = send.send_campaign(path, where["contacts"], where["out"], senders=senders, followup=True,
                                 per_company=per_company, suppression_path=where["suppression"],
                                 say=lambda *a: None)
        info["followup"] = {"days": int(campaign["followup"].get("days") or send.FOLLOWUP_DAYS),
                            "due": due["pending"], "this_run": due["this_run"], "waiting": due["followup_waiting"],
                            "sent": counts.get(send.FOLLOWUP, 0)}
    return info


def _followup_text(campaign):
    """The follow-up as the editor shows it, with a preview of how it arrives."""
    f = campaign.get("followup") or {}
    body = f.get("body") or []
    body = [body] if isinstance(body, str) else body
    sender = {"email": campaign["sender"]["email"], "name": campaign["sender"].get("name", "")}
    sample = {"email": "prospecto@ejemplo.com", "name": "Juan Pérez", "greeting": "Hola Juan,"}
    preview = ""
    if body:
        msg = send.make_followup(campaign, sample, "", sender)
        preview = msg.get_body(("plain",)).get_content().split("\n\nPara no recibir")[0]
        preview = f"Asunto: {msg['Subject']}\n\n{preview}"
    return {"days": int(f.get("days") or send.FOLLOWUP_DAYS), "body": "\n\n".join(body),
            "firma": "\n".join(send.followup_signature(campaign, sender)), "preview": preview,
            "editable": "_folder" in campaign}


def _reply_counts(out):
    counts = {}
    for row in replies.read_replies(out).values():
        counts[row["tipo"]] = counts.get(row["tipo"], 0) + 1
    return counts


EXPORTS = {
    "enviados": ["fecha", "email", "nombre", "empresa", "puesto", "remitente"],
    "pendientes": ["email", "nombre", "empresa", "puesto", "saludo"],
    "respuestas": ["fecha", "email", "nombre", "empresa", "puesto", "tipo", "regresa", "nuevo_contacto", "asunto"],
}


def export_csv(out_dir, slug, which):
    """CSV (Excel-ready) of who already got the campaign, or who is still pending.
    Pending uses the same list the next send would use (no unsubscribes or bounces)."""
    if which not in EXPORTS:
        raise ValueError("Elige enviados, pendientes o respuestas.")
    path = campaign_path(out_dir, slug)
    campaign = load_campaign(path)
    where = paths(out_dir, campaign)
    sent = {}
    for row in send.read_log(where["out"] / send.LOG_NAME):
        if row["estado"] == send.SENT:
            sent.setdefault(row["email"].lower(), row)
    contacts = {}
    if where["contacts"].exists():
        with open(where["contacts"], newline="", encoding="utf-8-sig") as fh:
            contacts = {(r.get("email") or "").strip().lower(): r for r in csv.DictReader(fh)}
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(EXPORTS[which])
    if which == "enviados":
        for email_addr, row in sent.items():
            c = contacts.get(email_addr, {})
            w.writerow([row["fecha"], email_addr, c.get("name", ""), c.get("company", ""), c.get("title", ""),
                        row.get("remitente", "")])
    elif which == "respuestas":  # the leads first
        order = {replies.REPLIED: 0, replies.GONE: 1, replies.AWAY: 2}
        for row in sorted(replies.read_replies(where["out"]).values(),
                          key=lambda r: (order.get(r["tipo"], 3), r["fecha"])):
            c = contacts.get(row["email"], {})
            w.writerow([row["fecha"], row["email"], c.get("name", ""), c.get("company", ""), c.get("title", ""),
                        row["tipo"], row["regresa"], row["nuevo_contacto"], row["asunto"]])
    elif where["contacts"].exists():
        people, _ = send.load_people(where["contacts"], where["out"], campaign, where["suppression"])
        for p in people:
            if p["email"] not in sent:
                w.writerow([p["email"], p["name"], p["company"], p["title"], p["greeting"]])
    return "\ufeff" + out.getvalue(), f"{path.stem}-{which}.csv"


def check_unsubscribes(out_dir, env_path, slug="", progress=None, stop=None):
    """Read the inboxes for 'Baja de correos' replies and bounces. Returns (new unsubs, new bounces)."""
    from mail_contact_scraper.config import load_settings
    from mail_contact_scraper.imap_source import connect as imap_connect

    campaign = load_campaign(campaign_path(out_dir, slug))
    boxes = send.inboxes(load_settings(env_path))
    if not boxes:
        raise ValueError("No hay buzones con contraseña para revisar.")
    own = campaign["sender"]["email"].split("@")[-1]
    try:
        where = paths(out_dir, campaign)
        return send.sync_suppression(boxes, where["suppression"], imap_connect, own_domains=[own],
                                     seen_path=where["suppression"].parent / send.SEEN_NAME, progress=progress, stop=stop)
    except send.Stopped:
        raise
    except Exception as exc:  # IMAP errors come in many shapes
        raise ValueError(f"No pude revisar bajas y rebotes: {exc}") from None


def check_replies(out_dir, env_path, slug="", progress=None, stop=None):
    """Read the inboxes for answers to the campaign and its follow-up. Returns {tipo: [emails]}."""
    from mail_contact_scraper.config import load_settings
    from mail_contact_scraper.imap_source import connect as imap_connect

    campaign = load_campaign(campaign_path(out_dir, slug))
    where = paths(out_dir, campaign)
    boxes = send.inboxes(load_settings(env_path))
    if not boxes:
        raise ValueError("No hay buzones con contraseña para revisar.")
    sent = [(r["email"], r.get("detalle", "")) for r in send.read_log(where["out"] / send.LOG_NAME)
            if r["estado"] in (send.SENT, send.FOLLOWUP)]
    if not sent:
        return {}
    own = campaign["sender"]["email"].split("@")[-1]
    try:
        return replies.sync_replies(boxes, imap_connect, where["out"], sent, where["suppression"], own_domains=[own],
                                    seen_path=where["out"] / send.SEEN_NAME, progress=progress, stop=stop)
    except send.Stopped:
        raise
    except Exception as exc:  # IMAP errors come in many shapes
        raise ValueError(f"No pude revisar las respuestas: {exc}") from None


def health(out_dir, env_path, slug=""):
    """SPF / DKIM / DMARC light of each sending domain the Campaña tab can pick."""
    campaign = load_campaign(campaign_path(out_dir, slug))
    by_domain = {}
    for address in available_senders(env_path, campaign):
        by_domain.setdefault(address.rsplit("@", 1)[-1].lower(), []).append(address)
    return [dict(deliverability.domain_health(d), senders=sorted(a)) for d, a in sorted(by_domain.items())]


def clean_list(out_dir, slug="", apply=False):
    """Pending addresses that would bounce for sure (domain with no mail server, typo in the
    domain). With apply=True they go to suppression.txt so no run sends to them."""
    path = campaign_path(out_dir, slug)
    campaign = load_campaign(path)
    where = paths(out_dir, campaign)
    if not where["contacts"].exists():
        raise ValueError("Todavía no hay contactos. Corre un escaneo primero.")
    people, _ = send.load_people(where["contacts"], where["out"], campaign, where["suppression"])
    done = {r["email"].lower() for r in send.read_log(where["out"] / send.LOG_NAME) if r["estado"] == send.SENT}
    found = deliverability.clean_list([p["email"] for p in people if p["email"] not in done])
    found["applied"] = 0
    if apply and found["bad"]:
        for reason, label in (("error de dedo", "error-de-dedo"), ("no recibe", "dominio-sin-correo")):
            group = [b["email"] for b in found["bad"] if reason in b["motivo"]]
            if group:
                send.add_suppression(where["suppression"], group, label)
                found["applied"] += len(group)
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["email", "motivo", "sugerencia"])
    for b in found["bad"]:
        w.writerow([b["email"], b["motivo"], b["sugerencia"]])
    found["csv"] = "\ufeff" + out.getvalue()
    return found


def _emails_in_file(filename, data):
    """Every address in an uploaded Excel or CSV, whatever column it is in."""
    name = (filename or "").lower()
    cells = []
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise ValueError("No pude abrir el Excel.") from None
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells += [str(v) for v in row if v]
    elif name.endswith((".csv", ".txt")):
        cells = [data.decode("utf-8-sig", errors="replace")]
    else:
        raise ValueError("Sube un Excel (.xlsx) o un CSV.")
    return {m.lower() for c in cells for m in re.findall(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", c)}


def keep_only(out_dir, env_path, slug, filename, data, apply=False):
    """Pending prospects of the campaign that are NOT in the uploaded list (the person cleaned
    it up: the rest are competitors). apply=True adds them to the suppression list and to the
    scraper's competitors list (a whole domain when every contact of it is in this group)."""
    from mail_contact_scraper.config import load_settings
    from mail_contact_scraper.contacts import COMPETITORS_HEADER, update_ignore_file

    keep = _emails_in_file(filename, data)
    if not keep:
        raise ValueError("No encontré correos en el archivo.")
    campaign = load_campaign(campaign_path(out_dir, slug))
    where = paths(out_dir, campaign)
    if not where["contacts"].exists():
        raise ValueError("Todavía no hay contactos. Corre un escaneo primero.")
    people, _ = send.load_people(where["contacts"], where["out"], campaign, where["suppression"])
    sent = {r["email"].lower() for r in send.read_log(where["out"] / send.LOG_NAME) if r["estado"] == send.SENT}
    pending = [p for p in people if p["email"] not in sent]
    extra = [p for p in pending if p["email"] not in keep]
    if len(extra) == len(pending):
        raise ValueError("Ningún correo del archivo está en los pendientes de esta campaña; revisa que sea la lista correcta.")
    with open(where["contacts"], newline="", encoding="utf-8-sig") as fh:
        everyone = [(r.get("email") or "").strip().lower() for r in csv.DictReader(fh)]
    extra_emails = {p["email"] for p in extra}
    entries = []
    for p in extra:
        domain = p["email"].rsplit("@", 1)[-1]
        whole = domain not in deliverability.KNOWN_DOMAINS and all(
            e in extra_emails for e in everyone if e.rsplit("@", 1)[-1] == domain)
        entry = domain if whole else p["email"]
        if entry not in entries:
            entries.append(entry)
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["email", "nombre", "empresa", "puesto"])
    for p in extra:
        w.writerow([p["email"], p["name"], p["company"], p["title"]])
    result = {"pending": len(pending), "in_file": len(pending) - len(extra), "extra": len(extra),
              "entries": entries, "csv": "\ufeff" + out.getvalue(), "applied": False}
    if apply and extra:
        send.add_suppression(where["suppression"], extra_emails, "competencia")
        update_ignore_file(load_settings(env_path).competitors_file, add=" ".join(entries), header=COMPETITORS_HEADER)
        result["applied"] = True
    return result


def bounce_checkpoint(check, out_dir, env_path, slug):
    """Hook for send_campaign: reads the inboxes and pauses when too many of this run bounce."""
    bounced_here = set()

    def checkpoint(sent_emails):
        try:
            _, bounced = check(out_dir, env_path, slug)
        except ValueError:
            return ""  # can't read the inbox now; the in-run breaker still watches
        sent = {e.lower() for e in sent_emails}
        bounced_here.update(b.lower() for b in bounced if b.lower() in sent)
        rate = len(bounced_here) / max(1, len(sent))
        if rate > send.BREAKER_RATE:
            return (f"{len(bounced_here)} de {len(sent)} correos de esta corrida rebotaron ({rate:.0%}); "
                    f"el máximo sano es {send.BREAKER_RATE:.0%}. Usa Limpiar lista antes de seguir.")
        return ""
    return checkpoint


class SendRun:
    """One campaign run at a time, in a background thread of the app."""

    def __init__(self, check=check_unsubscribes, answers=None):
        self._lock = threading.Lock()
        self._stop = threading.Event()
        self._thread = None
        self._check = check
        self._answers = answers or (lambda *a: check_replies(*a))
        self.state = {"running": False}

    def status(self):
        with self._lock:
            return dict(self.state)

    def _set(self, **values):
        with self._lock:
            self.state.update(values)

    def stop(self):
        if not self.status().get("running"):
            return False
        self._stop.set()
        self._set(phase="deteniendo")
        return True

    def start(self, out_dir, env_path, body, pause=4.0):
        if (body.get("confirm") or "").strip() != CONFIRM_WORD:
            raise ValueError(f"Escribe {CONFIRM_WORD} para confirmar el envío.")
        if self.status().get("running"):
            raise ValueError("Ya hay un envío en curso.")
        path = campaign_path(out_dir, str(body.get("campaign") or ""))
        campaign = load_campaign(path)
        where = paths(out_dir, campaign)
        if not where["contacts"].exists():
            raise ValueError("Todavía no hay contactos. Corre un escaneo primero.")
        followup = body.get("mode") == "seguimiento"
        if followup and not campaign.get("followup", {}).get("body"):
            raise ValueError("Esta campaña no tiene texto de seguimiento.")
        senders, _ = _senders(env_path, campaign, body.get("from"), _limit(body.get("daily_limit"), "El límite diario"))
        limit = _limit(body.get("limit"), "El máximo de esta corrida")
        options = {"limit": limit, "pause": pause, "followup": followup,
                   "per_company": _per_company(body.get("per_company"))}
        # Skipping the inbox check only skips reading new opt-outs and bounces; the ones already
        # in the suppression list still apply. A follow-up always checks: it must know who answered.
        check_first = followup or not body.get("skip_check")
        self._stop.clear()
        with self._lock:
            self.state = {"running": True, "phase": "bajas" if check_first else "enviando", "campaign": path.stem, "mode": body.get("mode") or "",
                          "senders": [s["email"] for s in senders], "started_at": _now(),
                          "sent": 0, "total": 0, "last": "", "note": ""}
        self._thread = threading.Thread(target=self._run, daemon=True,
                                        args=(out_dir, env_path, path, where, senders, options, check_first))
        self._thread.start()
        return self.status()

    def _run(self, out_dir, env_path, path, where, senders, options, check_first=True):
        note, awake = "", _keep_awake()
        try:
            if not check_first:
                note = "No revisé bajas ni rebotes antes de enviar; usé la lista que ya había."
            else:
                try:
                    note = self._check_first(out_dir, env_path, path, options)
                except send.Stopped:
                    self._set(note="Detenido mientras revisaba los buzones. No se envió nada.")
                    return
                if note is None:
                    return
            if self._stop.is_set():
                self._set(note=note)
                return
            self._set(phase="enviando", note=note, checking="")

            def progress(summary, email_addr, who):
                self._set(sent=summary["sent"], total=summary["this_run"], last=f"{email_addr} ({who})")

            result = send.send_campaign(
                path, where["contacts"], where["out"], senders=senders, send=True, **options,
                suppression_path=where["suppression"], sleep=self._stop.wait, stop=self._stop.is_set,
                progress=progress, say=lambda *a: None,
                checkpoint=bounce_checkpoint(self._check, out_dir, env_path, path.stem))
            self._set(sent=result["sent"], total=result["this_run"], rejected=len(result["rejected"]),
                      errors=result["errors"], stopped=result["stopped"], paused=result["paused"])
            if result["paused"]:
                notify("Envío en pausa", result["paused"])
            else:
                notify("Envío terminado", f"{result['sent']} enviados.")
        except Exception as exc:  # keep the app alive and show what happened
            self._set(error=str(exc) or exc.__class__.__name__)
            notify("El envío se detuvo", str(exc) or exc.__class__.__name__)
        finally:
            if awake:
                awake.terminate()
            self._set(running=False, phase="terminado", finished_at=_now())

    def _check_first(self, out_dir, env_path, path, options):
        """Read new opt-outs, bounces and replies. Returns the note, or None when the run must not go on."""
        progress = lambda text: self._set(checking=text)
        try:
            unsub, bounced = self._check(out_dir, env_path, path.stem, progress, self._stop.is_set)
            note = f"Bajas nuevas: {len(unsub)} · rebotes o bloqueos nuevos: {len(bounced)}."
        except ValueError as exc:
            note = f"{exc} Sigo con la lista que ya había."
        try:
            note += " " + _answers_note(self._answers(out_dir, env_path, path.stem, progress, self._stop.is_set))
        except ValueError as exc:
            if options["followup"]:  # without the replies we could write to someone who answered
                self._set(note=note, error=f"{exc} No mandé el seguimiento.")
                return None
            note += f" {exc}"
        return note


def _answers_note(found):
    parts = [f"{len(found.get(replies.REPLIED, []))} respuestas nuevas"]
    if found.get(replies.AWAY):
        parts.append(f"{len(found[replies.AWAY])} fuera de oficina")
    if found.get(replies.GONE):
        parts.append(f"{len(found[replies.GONE])} ya no laboran ahí")
    return " · ".join(parts).capitalize() + "."


def _keep_awake():
    """On a Mac, keep it from going to sleep while a run lasts (it can take hours)."""
    if sys.platform != "darwin":
        return None
    try:
        return subprocess.Popen(["/usr/bin/caffeinate", "-i"])
    except OSError:
        return None


def notify(title, text):
    """A macOS notification, so a pause is seen even with the app in the background."""
    if sys.platform != "darwin":
        return
    quote = lambda v: '"' + str(v).replace("\\", "\\\\").replace('"', '\\"')[:240] + '"'
    try:
        subprocess.Popen(["/usr/bin/osascript", "-e",
                          f"display notification {quote(text)} with title {quote(title)} sound name \"Basso\""])
    except OSError:
        pass


def _now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")
